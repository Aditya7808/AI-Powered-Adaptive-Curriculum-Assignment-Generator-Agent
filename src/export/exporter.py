import os

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

from src.schemas.assignment import Assignment
from src.schemas.curriculum import Curriculum


def export_curriculum_markdown(curriculum: Curriculum) -> str:
    """
    Exports a Curriculum object to human-readable Markdown format.
    """
    lines = [
        f"# {curriculum.title}",
        f"**Curriculum ID**: {curriculum.curriculum_id} | **Learner ID**: {curriculum.learner_id} | **Version**: v{curriculum.version}",
        f"**Created At**: {curriculum.created_at.strftime('%Y-%m-%d %H:%M:%S UTC')}",
        "",
        "## Alignment Notes",
        curriculum.alignment_notes,
        "",
        "## Modules Overview",
        "",
    ]

    for i, mod in enumerate(curriculum.modules, start=1):
        lines.append(f"### Module {i}: {mod.title} ({mod.week_range})")
        lines.append(f"- **Difficulty**: {mod.difficulty.value.title()}")
        lines.append(f"- **Target Competencies**: {', '.join(mod.competencies)}")
        lines.append("- **Learning Outcomes**:")
        for lo in mod.learning_outcomes:
            lines.append(f"  * {lo}")
        lines.append("- **Topics Covered**:")
        for t in mod.topics:
            lines.append(f"  * {t}")
        lines.append("- **Activities & Projects**:")
        for a in mod.activities:
            lines.append(f"  * {a}")
        if mod.resources:
            lines.append("- **Referenced Resources**:")
            for r in mod.resources:
                lines.append(f"  * [{r.file_name} p.{r.page}] {r.title}: _{r.snippet}_")
        lines.append("")

    return "\n".join(lines)


def export_curriculum_json(curriculum: Curriculum) -> str:
    """
    Exports a Curriculum object to formatted JSON.
    """
    return curriculum.model_dump_json(indent=2)


def export_curriculum_docx(curriculum: Curriculum, output_path: str) -> str:
    """
    Exports a Curriculum object to a professionally formatted Word (.docx) document.
    """
    os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)
    doc = Document()

    # Title
    title_p = doc.add_heading(curriculum.title, level=0)
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Subtitle / Metadata
    meta_p = doc.add_paragraph()
    meta_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta_p.add_run(
        f"Version {curriculum.version} | Learner: {curriculum.learner_id}\n"
    ).italic = True
    meta_p.add_run(f"Generated: {curriculum.created_at.strftime('%B %d, %Y')}")

    doc.add_heading("Curriculum Alignment & Adaptation Notes", level=1)
    doc.add_paragraph(curriculum.alignment_notes)

    doc.add_heading("Course Modules", level=1)

    for i, mod in enumerate(curriculum.modules, start=1):
        doc.add_heading(f"Module {i}: {mod.title}", level=2)

        # Meta info
        p = doc.add_paragraph()
        p.add_run("Timeline: ").bold = True
        p.add_run(f"{mod.week_range}  |  ")
        p.add_run("Difficulty: ").bold = True
        p.add_run(f"{mod.difficulty.value.title()}  |  ")
        p.add_run("Competencies: ").bold = True
        p.add_run(f"{', '.join(mod.competencies)}")

        # Learning Outcomes
        doc.add_heading("Learning Outcomes", level=3)
        for lo in mod.learning_outcomes:
            doc.add_paragraph(lo, style="List Bullet")

        # Topics
        doc.add_heading("Topics Covered", level=3)
        for topic in mod.topics:
            doc.add_paragraph(topic, style="List Bullet")

        # Activities
        doc.add_heading("Activities & Assignments", level=3)
        for act in mod.activities:
            doc.add_paragraph(act, style="List Bullet")

        # Resources
        if mod.resources:
            doc.add_heading("Grounding Resources", level=3)
            for res in mod.resources:
                doc.add_paragraph(
                    f"{res.title} ({res.file_name}, page {res.page}): {res.snippet}",
                    style="List Bullet",
                )

    doc.save(output_path)
    return output_path


def export_assignment_markdown(
    assignment: Assignment, include_solutions: bool = False
) -> str:
    """
    Exports an Assignment object to Markdown.
    If include_solutions=True, includes answer keys, explanations, and rubrics.
    """
    lines = [
        f"# {assignment.title}",
        f"**Assignment ID**: {assignment.assignment_id} | **Difficulty**: {assignment.difficulty.value.title()}",
        f"**Total Marks**: {assignment.total_marks} | **Estimated Time**: {assignment.estimated_time_minutes} minutes",
        "",
        "## Instructions",
        assignment.instructions,
        "",
        "## Learning Outcomes Assessed",
        "\n".join(f"- {lo}" for lo in assignment.learning_outcomes),
        "",
        "## Questions",
        "",
    ]

    for i, q in enumerate(assignment.questions, start=1):
        lines.append(f"### Question {i} ({q.marks} Marks)")
        lines.append(
            f"- **Competency**: {q.competency} | **Bloom Level**: {q.bloom_level.value.title()}"
        )
        lines.append(f"- **Type**: {q.type.value}")
        lines.append("")

        if q.scenario_context:
            lines.append(f"> **Scenario Context**:\n> {q.scenario_context}\n")

        lines.append(f"**Prompt**: {q.text}")
        lines.append("")

        if q.options:
            lines.append("**Options**:")
            for opt in q.options:
                lines.append(f"- [ ] {opt}")
            lines.append("")

        if include_solutions:
            lines.append("#### Solution & Answer Key:")
            lines.append(f"- **Correct Answer**: `{q.answer_key.correct_answer}`")
            lines.append(f"- **Explanation**: {q.answer_key.explanation}")
            if q.answer_key.common_mistakes:
                lines.append("- **Common Mistakes**:")
                for m in q.answer_key.common_mistakes:
                    lines.append(f"  * {m}")
            lines.append("")

    if include_solutions and assignment.rubric:
        lines.append("## Grading Rubric")
        for crit in assignment.rubric:
            lines.append(f"### Criterion: {crit.criterion} (Max {crit.max_points} pts)")
            lines.append(f"_{crit.description}_")
            for lvl in crit.levels:
                lines.append(f"- **{lvl.label} ({lvl.points} pts)**: {lvl.descriptor}")
            lines.append("")

    return "\n".join(lines)


def export_assignment_json(assignment: Assignment) -> str:
    """
    Exports an Assignment object to formatted JSON.
    """
    return assignment.model_dump_json(indent=2)


def export_assignment_docx(
    assignment: Assignment, output_path: str, include_solutions: bool = False
) -> str:
    """
    Exports an Assignment to a formatted Word (.docx) document.
    """
    os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)
    doc = Document()

    title_p = doc.add_heading(assignment.title, level=0)
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    meta_p = doc.add_paragraph()
    meta_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta_p.add_run(
        f"Total Marks: {assignment.total_marks}  |  Time: {assignment.estimated_time_minutes} mins  |  Difficulty: {assignment.difficulty.value.title()}"
    )

    doc.add_heading("Instructions", level=1)
    doc.add_paragraph(assignment.instructions)

    doc.add_heading("Questions", level=1)

    for i, q in enumerate(assignment.questions, start=1):
        doc.add_heading(f"Question {i} [{q.marks} Marks]", level=2)
        meta = doc.add_paragraph()
        meta.add_run(
            f"Competency: {q.competency} | Bloom Level: {q.bloom_level.value.title()} | Type: {q.type.value}"
        ).italic = True

        if q.scenario_context:
            sc_p = doc.add_paragraph()
            sc_p.add_run(f"Scenario:\n{q.scenario_context}").italic = True

        doc.add_paragraph(q.text)

        if q.options:
            for opt in q.options:
                doc.add_paragraph(f"[   ] {opt}", style="List Bullet")

        if include_solutions:
            sol_heading = doc.add_heading("Solution & Explanation", level=3)
            doc.add_paragraph(
                f"Correct Answer: {q.answer_key.correct_answer}"
            ).bold = True
            doc.add_paragraph(f"Explanation: {q.answer_key.explanation}")
            if q.answer_key.common_mistakes:
                doc.add_paragraph(
                    "Common Mistakes: " + ", ".join(q.answer_key.common_mistakes)
                )

    if include_solutions and assignment.rubric:
        doc.add_heading("Evaluation Rubric", level=1)
        for crit in assignment.rubric:
            doc.add_heading(f"{crit.criterion} (Max {crit.max_points} Points)", level=2)
            doc.add_paragraph(crit.description)
            table = doc.add_table(rows=1, cols=3)
            hdr_cells = table.rows[0].cells
            hdr_cells[0].text = "Level"
            hdr_cells[1].text = "Points"
            hdr_cells[2].text = "Descriptor"
            for lvl in crit.levels:
                row_cells = table.add_row().cells
                row_cells[0].text = lvl.label
                row_cells[1].text = str(lvl.points)
                row_cells[2].text = lvl.descriptor

    doc.save(output_path)
    return output_path
