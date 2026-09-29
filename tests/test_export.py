import json
import os
from datetime import datetime, timezone

import docx
import pytest

from src.export.exporter import (
    export_assignment_docx,
    export_assignment_json,
    export_assignment_markdown,
    export_curriculum_docx,
    export_curriculum_json,
    export_curriculum_markdown,
)
from src.schemas.assignment import (
    AnswerKey,
    Assignment,
    Question,
    RubricCriterion,
    RubricLevel,
)
from src.schemas.curriculum import Curriculum, Module, ResourceReference
from src.schemas.learner import BloomLevel, DifficultyLevel, QuestionType


@pytest.fixture
def sample_curriculum():
    mods = [
        Module(
            module_id=f"mod-{i}",
            title=f"Module {i}: Foundations of AI",
            week_range=f"Week {i}",
            competencies=["Machine Learning", "Python"],
            learning_outcomes=["Understand ML models"],
            topics=["Supervised Learning", "Linear Regression"],
            activities=["Hands-on lab 1"],
            resources=[
                ResourceReference(
                    title="ML Primer",
                    file_name="ml_book.pdf",
                    page=12,
                    chunk_id="chk-12",
                    snippet="Linear regression minimizes MSE.",
                )
            ],
            prerequisites=[],
            difficulty=DifficultyLevel.medium,
        )
        for i in range(1, 4)
    ]
    return Curriculum(
        curriculum_id="curr-exp-1",
        learner_id="learner-exp-1",
        version=1,
        title="Applied Machine Learning",
        modules=mods,
        alignment_notes="Designed for intermediate python developers.",
        created_at=datetime.now(timezone.utc),
    )


@pytest.fixture
def sample_assignment():
    return Assignment(
        assignment_id="assign-exp-1",
        curriculum_id="curr-exp-1",
        module_id="mod-1",
        title="Module 1 Quiz: Linear Regression",
        difficulty=DifficultyLevel.medium,
        instructions="Answer the following questions within 30 minutes.",
        questions=[
            Question(
                question_id="q1",
                text="Which cost function is commonly used for linear regression?",
                type=QuestionType.mcq,
                difficulty=DifficultyLevel.easy,
                bloom_level=BloomLevel.remember,
                competency="Machine Learning",
                options=[
                    "Mean Squared Error",
                    "Cross Entropy",
                    "Hinge Loss",
                    "Dice Loss",
                ],
                answer_key=AnswerKey(
                    correct_answer="Mean Squared Error",
                    explanation="MSE calculates the average squared difference between predictions and actuals.",
                    common_mistakes=["Cross Entropy is for classification"],
                ),
                marks=5,
            ),
            Question(
                question_id="q2",
                text="Explain the effect of high multicollinearity on linear regression coefficients.",
                type=QuestionType.short_answer,
                difficulty=DifficultyLevel.medium,
                bloom_level=BloomLevel.understand,
                competency="Statistical Learning",
                options=None,
                answer_key=AnswerKey(
                    correct_answer="Multicollinearity inflates the variance of coefficient estimates, making them unstable.",
                    explanation="Coefficients become sensitive to minor changes in the model.",
                    common_mistakes=["Claiming predictions become impossible"],
                ),
                marks=5,
            ),
        ],
        rubric=[
            RubricCriterion(
                criterion="Statistical Rigor",
                description="Depth and accuracy of explanation",
                levels=[
                    RubricLevel(
                        label="Advanced", points=5, descriptor="Rigorous and clear"
                    ),
                    RubricLevel(
                        label="Basic", points=2, descriptor="Surface understanding"
                    ),
                ],
                max_points=5,
            )
        ],
        learning_outcomes=[
            "Identify appropriate loss functions",
            "Diagnose multicollinearity",
        ],
        total_marks=10,
        estimated_time_minutes=30,
        status="approved",
    )


def test_export_curriculum_markdown(sample_curriculum):
    md = export_curriculum_markdown(sample_curriculum)
    assert "# Applied Machine Learning" in md
    assert "Module 1: Foundations of AI" in md
    assert "Linear regression minimizes MSE." in md
    assert "Designed for intermediate python developers." in md


def test_export_curriculum_json(sample_curriculum):
    js = export_curriculum_json(sample_curriculum)
    data = json.loads(js)
    assert data["curriculum_id"] == "curr-exp-1"
    assert len(data["modules"]) == 3


def test_export_curriculum_docx(sample_curriculum, tmp_path):
    out_file = str(tmp_path / "curriculum.docx")
    path = export_curriculum_docx(sample_curriculum, out_file)
    assert os.path.exists(path)

    doc = docx.Document(path)
    text = " ".join(p.text for p in doc.paragraphs)
    assert "Applied Machine Learning" in text
    assert "Module 1: Foundations of AI" in text


def test_export_assignment_markdown(sample_assignment):
    # Without solutions
    md_student = export_assignment_markdown(sample_assignment, include_solutions=False)
    assert "Module 1 Quiz: Linear Regression" in md_student
    assert "Which cost function is commonly used for linear regression?" in md_student
    assert "Solution & Answer Key:" not in md_student

    # With solutions
    md_teacher = export_assignment_markdown(sample_assignment, include_solutions=True)
    assert "Solution & Answer Key:" in md_teacher
    assert "Mean Squared Error" in md_teacher
    assert "Grading Rubric" in md_teacher


def test_export_assignment_json(sample_assignment):
    js = export_assignment_json(sample_assignment)
    data = json.loads(js)
    assert data["assignment_id"] == "assign-exp-1"
    assert len(data["questions"]) == 2


def test_export_assignment_docx(sample_assignment, tmp_path):
    out_file = str(tmp_path / "assignment.docx")
    path = export_assignment_docx(sample_assignment, out_file, include_solutions=True)
    assert os.path.exists(path)

    doc = docx.Document(path)
    text = " ".join(p.text for p in doc.paragraphs)
    assert "Module 1 Quiz: Linear Regression" in text
    assert "Solution & Explanation" in text
    assert len(doc.tables) > 0  # Rubric table
